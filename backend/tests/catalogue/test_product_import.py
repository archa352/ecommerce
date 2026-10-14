"""Bulk product import from CSV."""

from __future__ import annotations

from tests.catalogue.factories import PRODUCTS

HEADER = "product_name,category,sku,variant_name,price,stock,size,colour,status,fabric,image_url\n"


def _csv(*rows: str) -> bytes:
    return (HEADER + "\n".join(rows) + "\n").encode()


async def _import(client, headers, data: bytes, dry_run: bool):
    return await client.post(
        f"{PRODUCTS}/import",
        files={"file": ("products.csv", data, "text/csv")},
        data={"dry_run": str(dry_run).lower()},
        headers=headers,
    )


async def test_a_dry_run_validates_without_creating(client, admin_headers, category):
    data = _csv(
        f"Cotton Kurta,{category['name']},CK-S,,1299,5,S,Indigo,ACTIVE,Cotton,https://cdn.example.com/a.jpg",
        f"Cotton Kurta,{category['name']},CK-M,,1299,3,M,Indigo,ACTIVE,Cotton,https://cdn.example.com/a.jpg",
    )

    response = await _import(client, admin_headers, data, dry_run=True)
    listed = await client.get(PRODUCTS, headers=admin_headers)

    body = response.json()["data"]
    assert body == {"rows": 2, "products": 1, "variants": 2, "dry_run": True, "errors": []}
    assert listed.json()["data"] == []


async def test_an_import_creates_products_and_variants(client, admin_headers, category):
    data = _csv(
        f"Rayon Set,{category['name']},RS-S,,999,4,S,Red,ACTIVE,Rayon,https://cdn.example.com/b.jpg",
        f"Rayon Set,{category['name']},RS-M,,999,6,M,Red,ACTIVE,Rayon,https://cdn.example.com/b.jpg",
    )

    response = await _import(client, admin_headers, data, dry_run=False)
    products = (await client.get(PRODUCTS, headers=admin_headers)).json()["data"]
    variants = (await client.get(f"{PRODUCTS}/{products[0]['id']}/variants", headers=admin_headers)).json()["data"]

    assert response.json()["message"] == "Products imported"
    assert products[0]["fabric"] == "Rayon"
    assert {v["sku"] for v in variants} == {"RS-S", "RS-M"}
    assert {v["options"]["Size"] for v in variants} == {"S", "M"}
    assert sum(v["inventory"]["available_quantity"] for v in variants) == 10


async def test_bad_rows_block_the_whole_file(client, admin_headers, category):
    data = _csv(
        f"Good,{category['name']},G-1,,100,1,,,,,https://cdn.example.com/g.jpg",
        "Bad,No Such Category,B-1,,100,1,,,,,https://cdn.example.com/g.jpg",
        f"Bad Price,{category['name']},B-2,,-5,1,,,,,https://cdn.example.com/g.jpg",
    )

    response = await _import(client, admin_headers, data, dry_run=False)
    listed = await client.get(PRODUCTS, headers=admin_headers)

    errors = response.json()["data"]["errors"]
    assert [e["row"] for e in errors] == [3, 4]
    assert listed.json()["data"] == []


async def test_existing_skus_are_reported_not_overwritten(client, admin_headers, category, variant):
    data = _csv(f"Again,{category['name']},{variant['sku']},,100,1,,,,,https://cdn.example.com/g.jpg")

    response = await _import(client, admin_headers, data, dry_run=True)

    assert "already exists" in response.json()["data"]["errors"][0]["message"]


async def test_missing_columns_are_rejected(client, admin_headers):
    response = await _import(client, admin_headers, b"name,price\nx,1\n", dry_run=True)

    assert response.status_code == 400


async def test_import_needs_catalogue_manage(client, staff_headers):
    response = await _import(client, staff_headers, _csv(), dry_run=True)

    assert response.status_code == 403


async def test_the_template_has_the_columns(client, admin_headers):
    response = await client.get(f"{PRODUCTS}/import/template", headers=admin_headers)

    assert response.text.startswith("product_name,category,sku")


# ------------------------------------------------------------------ Excel (.xlsx)


def _xlsx(*rows: list) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADER.strip().split(","))
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _import_xlsx(client, headers, data: bytes, dry_run: bool, name="products.xlsx"):
    return await client.post(
        f"{PRODUCTS}/import",
        files={"file": (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"dry_run": str(dry_run).lower()},
        headers=headers,
    )


def _row(category, sku, size, stock=4, price=999):
    # Excel stores numbers as numbers -- the importer must read 4 / 999.0 like "4" / "999".
    return ["Excel Kurta", category["name"], sku, None, float(price), stock, size, "Blue", "ACTIVE", "Cotton", "https://cdn.example.com/x.jpg"]


async def test_an_excel_dry_run_validates_without_creating(client, admin_headers, category):
    data = _xlsx(_row(category, "XK-S", "S"), [], _row(category, "XK-M", "M"))

    response = await _import_xlsx(client, admin_headers, data, dry_run=True)
    listed = await client.get(PRODUCTS, headers=admin_headers)

    assert response.json()["data"] == {"rows": 2, "products": 1, "variants": 2, "dry_run": True, "errors": []}
    assert listed.json()["data"] == []


async def test_an_excel_import_creates_products_and_stock(client, admin_headers, category):
    response = await _import_xlsx(client, admin_headers, _xlsx(_row(category, "XK-S", "S", stock=7)), dry_run=False)
    products = (await client.get(PRODUCTS, headers=admin_headers)).json()["data"]
    variants = (await client.get(f"{PRODUCTS}/{products[0]['id']}/variants", headers=admin_headers)).json()["data"]

    assert response.json()["message"] == "Products imported"
    assert products[0]["slug"] == "excel-kurta"
    assert variants[0]["price"] == "999.00"
    assert variants[0]["inventory"]["available_quantity"] == 7


async def test_excel_rows_with_errors_are_reported_by_row(client, admin_headers, category):
    bad_category = ["Bad", "Nope", "B-1", None, 100, 1, None, None, None, None, "https://cdn.example.com/g.jpg"]
    bad_stock = _row(category, "B-2", "S", stock="lots")
    data = _xlsx(_row(category, "G-1", "S"), bad_category, bad_stock, _row(category, "G-1", "M"))

    errors = (await _import_xlsx(client, admin_headers, data, dry_run=False)).json()["data"]["errors"]
    listed = await client.get(PRODUCTS, headers=admin_headers)

    assert [e["row"] for e in errors] == [3, 4, 5]
    assert "appears twice" in errors[2]["message"]
    assert listed.json()["data"] == []


async def test_existing_skus_are_reported_in_excel_too(client, admin_headers, category, variant):
    errors = (await _import_xlsx(client, admin_headers, _xlsx(_row(category, variant["sku"], "S")), dry_run=True)).json()["data"]["errors"]

    assert "already exists" in errors[0]["message"]


async def test_an_excel_file_missing_columns_is_rejected(client, admin_headers):
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["name", "price"])
    buffer = BytesIO()
    workbook.save(buffer)

    response = await _import_xlsx(client, admin_headers, buffer.getvalue(), dry_run=True)

    assert response.status_code == 400
    assert "Missing columns" in response.json()["message"]


async def test_a_broken_excel_file_is_rejected(client, admin_headers):
    response = await _import_xlsx(client, admin_headers, b"not really a spreadsheet", dry_run=True)

    assert response.status_code == 400
    assert "not a valid .xlsx" in response.json()["message"]


async def test_other_file_types_are_rejected(client, admin_headers):
    response = await _import_xlsx(client, admin_headers, b"x", dry_run=True, name="products.xls")

    assert response.status_code == 400


async def test_the_excel_template_has_the_columns(client, admin_headers):
    from io import BytesIO

    from openpyxl import load_workbook

    response = await client.get(f"{PRODUCTS}/import/template", params={"format": "xlsx"}, headers=admin_headers)

    header = [c.value for c in next(load_workbook(BytesIO(response.content)).active.iter_rows(max_row=1))]
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert header[:3] == ["product_name", "category", "sku"]
