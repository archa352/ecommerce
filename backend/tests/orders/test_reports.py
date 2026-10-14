"""Admin reports and analytics, with CSV export."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

REPORTS = "/api/v1/admin/reports"
TODAY = date.today().isoformat()
RANGE = {"start_date": TODAY, "end_date": TODAY}


UNDATED = {"inventory", "unsold"}


async def _get(client, headers, name, **params):
    query = params if name in UNDATED else {**RANGE, **params}
    response = await client.get(f"{REPORTS}/{name}", params=query, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()["data"]


async def test_the_sales_report_buckets_and_totals(client, admin_headers, customer_auth, order_in):
    one = await order_in(customer_auth)
    two = await order_in(customer_auth, quantity=2)

    data = await _get(client, admin_headers, "sales", group_by="day")

    assert len(data["rows"]) == 1
    assert data["summary"]["orders"] == 2
    expected = Decimal(one["total_amount"]) + Decimal(two["total_amount"])
    assert Decimal(data["summary"]["revenue"]) == expected
    assert Decimal(data["summary"]["average_order_value"]) == (expected / 2).quantize(Decimal("0.01"))


async def test_cancelled_orders_are_excluded_from_sales(client, admin_headers, customer_auth, order_in):
    order = await order_in(customer_auth)
    await client.post(f"/api/v1/orders/{order['id']}/cancel", headers=customer_auth)

    data = await _get(client, admin_headers, "sales")

    assert data["summary"]["orders"] == 0
    assert data["summary"]["cancelled_orders"] == 1


async def test_sales_by_product_category_and_location(client, admin_headers, customer_auth, order_in):
    order = await order_in(customer_auth, quantity=3)

    products = await _get(client, admin_headers, "sales-by-product")
    categories = await _get(client, admin_headers, "sales-by-category")
    states = await _get(client, admin_headers, "sales-by-location")

    assert products[0]["units_sold"] == 3
    assert products[0]["product_id"] == order["items"][0]["product_id"]
    assert categories[0]["category_name"] == "Order Test Category"
    assert states[0]["state"] == "Karnataka"
    assert states[0]["orders"] == 1


async def test_the_inventory_report_flags_low_and_out_of_stock(client, admin_headers, new_variant):
    await new_variant(initial_quantity=0)
    await new_variant(initial_quantity=1, low_stock_threshold=3)
    await new_variant(initial_quantity=50)

    out = await _get(client, admin_headers, "inventory", stock_status="out")
    low = await _get(client, admin_headers, "inventory", stock_status="low")
    everything = await _get(client, admin_headers, "inventory")

    assert [r["stock_status"] for r in out] == ["OUT_OF_STOCK"]
    assert [r["sellable_quantity"] for r in low] == [1]
    assert len(everything) == 3


async def test_unsold_products_need_age(client, admin_headers, variant):
    rows = await _get(client, admin_headers, "unsold", days=30)

    assert rows == []  # listed today, so not "unsold for 30 days" yet


async def test_new_vs_returning_customers(client, admin_headers, customer_auth, order_in):
    await order_in(customer_auth)
    await order_in(customer_auth)

    data = await _get(client, admin_headers, "customers")

    assert data["new_customers"] == 1
    assert data["returning_customers"] == 0
    assert data["top_customers"][0]["orders"] == 2


async def test_the_coupon_report(client, admin_headers, customer_auth, variant, add_to_cart, address, place_order):
    await client.post(
        "/api/v1/admin/coupons",
        json={"code": "TENOFF", "discount_type": "FIXED_AMOUNT", "discount_value": "10"},
        headers=admin_headers,
    )
    await add_to_cart(customer_auth, variant["id"], 1)
    await place_order(customer_auth, address["id"], coupon_code="TENOFF")

    rows = await _get(client, admin_headers, "coupons")

    assert rows[0]["code"] == "TENOFF"
    assert rows[0]["orders"] == 1
    assert Decimal(rows[0]["discount_given"]) == Decimal("10.00")


async def test_the_returns_report(client, admin_headers, customer_auth, order_in, set_status):
    returned = await order_in(customer_auth)
    delivered = await order_in(customer_auth)
    for status in ("CONFIRMED", "PACKED", "SHIPPED"):
        await set_status(returned["id"], status)
        await set_status(delivered["id"], status)
    await set_status(returned["id"], "RETURNED")
    await set_status(delivered["id"], "DELIVERED")

    data = await _get(client, admin_headers, "returns")

    assert data["returned_orders"] == 1
    assert data["return_rate"] == 50.0
    assert data["by_product"][0]["units_returned"] == 1


async def test_analytics_kpis(client, admin_headers, customer_auth, order_in):
    order = await order_in(customer_auth)

    data = await _get(client, admin_headers, "analytics")

    assert data["orders"] == 1
    assert Decimal(data["average_order_value"]) == Decimal(order["total_amount"])
    assert data["cod_orders"] == 1
    assert data["prepaid_orders"] == 0
    assert data["cod_share_pct"] == 100.0


async def test_reports_export_as_csv(client, admin_headers, customer_auth, order_in):
    await order_in(customer_auth)

    response = await client.get(
        f"{REPORTS}/sales-by-product", params={**RANGE, "format": "csv"}, headers=admin_headers
    )

    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]
    assert response.text.lstrip("﻿").startswith("product_id,product_name,units_sold")


async def test_reports_need_permission(client, staff_headers):
    response = await client.get(f"{REPORTS}/sales", headers=staff_headers)

    assert response.status_code == 403


async def test_a_range_must_run_forwards(client, admin_headers):
    response = await client.get(
        f"{REPORTS}/sales", params={"start_date": "2026-02-01", "end_date": "2026-01-01"}, headers=admin_headers
    )

    assert response.status_code == 422


async def test_reports_are_per_tenant(other_client, other_admin_headers, customer_auth, order_in):
    await order_in(customer_auth)

    data = (await other_client.get(f"{REPORTS}/sales", params=RANGE, headers=other_admin_headers)).json()["data"]

    assert data["summary"]["orders"] == 0


# ------------------------------------------------------------ Excel / PDF downloads

EXPORTABLE = [
    ("sales", RANGE),
    ("sales-by-product", RANGE),
    ("sales-by-category", RANGE),
    ("sales-by-location", RANGE),
    ("inventory", {}),
    ("unsold", {}),
    ("customers", RANGE),
    ("coupons", RANGE),
    ("returns", RANGE),
    ("analytics", RANGE),
]


@pytest.mark.parametrize(("name", "params"), EXPORTABLE)
@pytest.mark.parametrize("fmt", ["csv", "xlsx", "pdf"])
async def test_every_report_downloads_in_every_format(client, admin_headers, name, params, fmt):
    response = await client.get(f"{REPORTS}/{name}", params={**params, "format": fmt}, headers=admin_headers)

    assert response.status_code == 200, response.text
    assert response.headers["content-disposition"].startswith("attachment; filename=")
    assert response.headers["content-disposition"].endswith(f'.{fmt}"')
    expected = {"csv": "text/csv", "xlsx": "application/vnd.openxmlformats", "pdf": "application/pdf"}[fmt]
    assert response.headers["content-type"].startswith(expected)


async def test_the_excel_file_has_headings_and_real_numbers(client, admin_headers, customer_auth, order_in):
    await order_in(customer_auth, quantity=3)

    response = await client.get(
        f"{REPORTS}/sales-by-product", params={**RANGE, "format": "xlsx"}, headers=admin_headers
    )
    json_rows = await _get(client, admin_headers, "sales-by-product")

    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:4] == ("Product id", "Product name", "Units sold", "Orders")
    assert rows[1][2] == 3  # a number, not text
    assert Decimal(str(rows[1][rows[0].index("Net sales")])) == Decimal(json_rows[0]["net_sales"])
    assert len(rows) == 1 + len(json_rows)


async def test_the_pdf_shows_the_title_period_and_rows(client, admin_headers, customer_auth, order_in):
    order = await order_in(customer_auth)

    response = await client.get(
        f"{REPORTS}/sales-by-product", params={**RANGE, "format": "pdf"}, headers=admin_headers
    )

    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert "Sales by product" in text
    assert f"{TODAY} to {TODAY}" in text
    assert order["items"][0]["product_name"] in text
    assert "Units sold" in text


async def test_an_empty_report_still_downloads(client, admin_headers):
    xlsx = await client.get(f"{REPORTS}/coupons", params={**RANGE, "format": "xlsx"}, headers=admin_headers)
    pdf = await client.get(f"{REPORTS}/coupons", params={**RANGE, "format": "pdf"}, headers=admin_headers)

    assert load_workbook(io.BytesIO(xlsx.content)).active["A1"].value == "No data for this report"
    assert "No data for this report" in PdfReader(io.BytesIO(pdf.content)).pages[0].extract_text()


async def test_analytics_downloads_product_performance(client, admin_headers, customer_auth, order_in):
    await order_in(customer_auth)

    response = await client.get(f"{REPORTS}/analytics", params={**RANGE, "format": "csv"}, headers=admin_headers)

    assert response.text.lstrip("﻿").startswith("product_id,")


async def test_an_unknown_format_is_rejected(client, admin_headers):
    response = await client.get(f"{REPORTS}/sales", params={"format": "docx"}, headers=admin_headers)

    assert response.status_code == 422
